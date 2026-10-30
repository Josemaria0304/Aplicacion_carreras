import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Style Header
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for i, title in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "0F172A") # Dark Navy
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Style Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)

    # Set Widths if provided
    if col_widths:
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)

    doc.add_paragraph() # Spacing
    return table

def add_callout(doc, title, text, tone="info"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.5)

    bg_hex = "F0FDF4" if tone == "success" else "F8FAFC"
    border_color = "16A34A" if tone == "success" else "0284C7"

    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(15, 23, 42)

    run_c = p.add_run(text)
    run_c.font.name = 'Calibri'
    run_c.font.size = Pt(9)
    run_c.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph()

def build_word_documentation(output_path):
    doc = Document()

    # Configure Margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.15

    # -------------------------------------------------------------
    # PORTADA (COVER PAGE)
    # -------------------------------------------------------------
    cover_p_space = doc.add_paragraph()
    cover_p_space.paragraph_format.space_before = Pt(36)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    run_cat = title_p.add_run("APEX RACE CONTROL PLATFORM\n")
    run_cat.font.name = 'Calibri'
    run_cat.font.size = Pt(12)
    run_cat.bold = True
    run_cat.font.color.rgb = RGBColor(2, 132, 199) # Sky blue

    run_title = title_p.add_run("Documentación Técnica y Manual de Operaciones CRUD")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(26)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42) # Dark Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(30)
    run_sub = sub_p.add_run("Módulo Completo de Inserción, Modificación, Eliminación y Actualización de Corredores, Autos y Catálogos de Competición con Sincronización MySQL y Fallback Local")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Table in Cover
    meta_headers = ["Elemento", "Detalle de Especificación"]
    meta_rows = [
        ["Proyecto", "Apex Race Control - GT Series 2026"],
        ["Versión de Software", "v2.0.0 (CRUD Engine & Unified Store Architecture)"],
        ["Módulos Cubiertos", "Pilotos (Drivers), Garaje / Vehículos (Vehicles), Equipos (Teams), Catálogos MySQL"],
        ["Stack Tecnológico", "React 19 + TypeScript 7 + Vite 8 + Tailwind CSS 4 + Zustand 5 + Express 5 + MySQL"],
        ["Fecha de Generación", "Octubre 2026"],
        ["Estado del Sistema", "Completado y validado para producción"],
    ]
    create_styled_table(doc, meta_headers, meta_rows, [2.2, 4.3])

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECCIÓN 1: RESUMEN EJECUTIVO
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Resumen Ejecutivo y Alcance del Proyecto", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Apex Race Control es una plataforma modular de alta gama diseñada para la administración integral, "
        "dirección deportiva, cronometraje en vivo y fiscalización técnica de campeonatos de automovilismo tipo Gran Turismo (GT3 / GT4). "
        "Originalmente concebida como una consola de control con datos estáticos y un simulador de telemetría de navegador, "
        "el proyecto ha sido evolucionado a una solución integral que permite gestionar activamente el ciclo de vida de los datos."
    )

    doc.add_paragraph(
        "Mediante la implementación de una arquitectura de almacenamiento reactivo y sincronizado, el sistema ahora ofrece:"
    )

    bullet_points = [
        ("Gestión Integral de Corredores (Pilotos): ", "Capacidad para inscribir nuevos corredores, actualizar fichas deportivas y biométricas, asignar escuderías y vehículos, y eliminar corredores con recálculo dinámico de posiciones."),
        ("Control Técnico de Autos (Vehículos / Parque Cerrado): ", "Registro de chasis de competición, configuración de motorizaciones, potencia (HP), torque (Nm), peso mínimo BoP, estado de Scrutineering (Aprobado, Pendiente, Observado, Rechazado) y transponders."),
        ("Persistencia Híbrida Inteligente: ", "Sincronización transparente con base de datos MySQL mediante API RESTful cuando el servidor está activo, y respaldo automático e instantáneo en LocalStorage si el entorno opera sin base de datos activa."),
        ("Eliminación Segura y Completa: ", "Adición de endpoints DELETE en el backend y modales interactivos de confirmación en la interfaz de usuario para prevenir pérdidas accidentales de información."),
        ("Administración de Catálogos Relacionales: ", "Panel de control avanzado con soporte para Equipos, Circuitos, Campeonatos, Eventos y Tablas de Puntuación.")
    ]

    for title, desc in bullet_points:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(title)
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        r_d = bp.add_run(desc)

    add_callout(
        doc,
        "Logro Clave de la Implementación",
        "Anteriormente, el garaje y los pilotos dependían de arreglos fijos de sólo lectura definidos en el archivo seed.ts. "
        "Con la nueva arquitectura, cualquier modificación, creación o baja realizada por el usuario se propaga de forma reactiva e inmediata "
        "a todas las vistas (Dashboard, Clasificaciones, Detalle de Pilotos, Garaje y Live Timing), persistiendo entre reinicios del navegador.",
        "success"
    )

    # -------------------------------------------------------------
    # SECCIÓN 2: ARQUITECTURA TÉCNICA DEL SISTEMA
    # -------------------------------------------------------------
    h1 = doc.add_heading("2. Arquitectura de Software y Capas del Sistema", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El sistema Apex Race Control se estructura siguiendo una arquitectura moderna por capas desacopladas, "
        "optimizando el rendimiento, la reactividad en el cliente y la integridad de los datos en el servidor:"
    )

    arch_headers = ["Capa de Software", "Tecnología Principal", "Función y Responsabilidad"]
    arch_data = [
        ["Frontend UI", "React 19 + TypeScript + Tailwind CSS 4", "Renderizado de interfaces modulares, tablas de cronometraje, modales de formulario y paneles de control."],
        ["Enrutamiento", "React Router 7", "Navegación SPA (Single Page Application) con rutas protegidas según roles de acceso."],
        ["Gestión de Estado Local", "Zustand 5 (useCompetition, useRace, useSession)", "Store central reactivo que gestiona los corredores, vehículos, clasificaciones y simulación."],
        ["Capa de Caché Cliente", "LocalStorage (apex_competition_store_v2)", "Almacenamiento persistente en navegador que garantiza disponibilidad inmediata y soporte offline."],
        ["Backend REST API", "Node.js + Express 5", "Servidor HTTP ligero encargado de recibir peticiones CRUD, validar esquemas y procesar transacciones."],
        ["Conectividad MySQL", "mysql2/promise con Connection Pool", "Gestión eficiente de conexiones a la base de datos relacional MySQL."],
        ["Base de Datos", "MySQL (esquema apex_race_control)", "Almacén de datos con soporte de tipos JSON nativos para conservar estructuras enriquecidas."]
    ]
    create_styled_table(doc, arch_headers, arch_data, [1.5, 2.0, 3.0])

    doc.add_paragraph(
        "A continuación se describe el flujo de datos unidireccional y bidireccional:\n"
        "1. El usuario interactúa con la interfaz gráfica (por ejemplo, presiona '+ Inscribir Piloto' o '+ Registrar Vehículo').\n"
        "2. El modal correspondiente valida los tipos de datos y envía el payload a useCompetition.\n"
        "3. El store aplica una actualización optimista en memoria y guarda una copia de seguridad en LocalStorage.\n"
        "4. En paralelo, si la API de MySQL está en línea, se emite una petición asíncrona (POST, PUT o DELETE) hacia /api/:resource.\n"
        "5. El servidor Express valida las reglas de negocio, ejecuta la sentencia SQL en MySQL y responde con el estado de la operación."
    )

    # -------------------------------------------------------------
    # SECCIÓN 3: ESTRUCTURA DE ARCHIVOS Y DIRECTORIOS
    # -------------------------------------------------------------
    h1 = doc.add_heading("3. Estructura del Código Fuente", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El proyecto mantiene una estructura modular limpia y organizada bajo el directorio raíz:"
    )

    files_headers = ["Directorio / Archivo", "Propósito en el Sistema"]
    files_data = [
        ["server/index.js", "Servidor Express con endpoints GET, POST, PUT, DELETE, /counts, /health y bootstrap."],
        ["database/schema.sql", "Definición del esquema DDL de MySQL con las 8 tablas de competición."],
        ["src/store/useCompetition.ts", "Store maestro reactivo con los métodos CRUD: addDriver, updateDriver, deleteDriver, addVehicle, updateVehicle, deleteVehicle, syncFromApi."],
        ["src/store/useRace.ts", "Motor de simulación matemática de la carrera en vivo (laps, telemetry, sectores, safety car)."],
        ["src/store/useSession.ts", "Store de autenticación simulada por roles (organizer, steward, timekeeper, pilot) y notificaciones toast."],
        ["src/components/DriverModal.tsx", "Componente modal interactivo para crear y modificar corredores con validación completa."],
        ["src/components/VehicleModal.tsx", "Componente modal interactivo para registrar y modificar vehículos y especificaciones técnicas."],
        ["src/components/ConfirmDeleteModal.tsx", "Componente modal de advertencia y confirmación segura de borrado."],
        ["src/components/DatabaseManager.tsx", "Consola administrativa con selector de catálogo, editor JSON y soporte de borrado directo."],
        ["src/pages/Drivers.tsx", "Página de visualización, búsqueda, filtrado y acciones rápidas de edición/borrado de corredores."],
        ["src/pages/DriverDetail.tsx", "Vista detallada de la ficha deportiva del corredor con botones de acción directa (Editar/Eliminar)."],
        ["src/pages/Vehicles.tsx", "Página de garaje con métricas técnicas (potencia total, aprobados) y tarjetas interactivas de autos."],
        ["src/pages/VehicleDetail.tsx", "Vista técnica de chasis con inspección Scrutineering y gestión de ciclo de vida del vehículo."],
        ["src/pages/Standings.tsx", "Tabla de posiciones del campeonato reactiva a las altas y bajas de pilotos y escuderías."],
        ["src/lib/data.ts", "Adaptador desacoplado que expone getters dinámicos conectados al store reactivo."]
    ]
    create_styled_table(doc, files_headers, files_data, [2.2, 4.3])

    # -------------------------------------------------------------
    # SECCIÓN 4: MODELO DE DATOS Y BASE DE DATOS MYSQL
    # -------------------------------------------------------------
    h1 = doc.add_heading("4. Esquema de Base de Datos y Modelo de Entidades", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "La base de datos se denomina apex_race_control con codificación utf8mb4_unicode_ci. "
        "Utiliza una arquitectura híbrida relacional-documental: cada tabla posee una clave primaria id (VARCHAR(64)), "
        "un campo data (JSON) para almacenar la entidad con todos sus atributos enriquecidos (incluyendo arrays y objetos anidados), "
        "y marcas de tiempo automáticas created_at y updated_at."
    )

    db_headers = ["Tabla en MySQL", "Entidad de Negocio", "Estructura Principal del Campo JSON"]
    db_data = [
        ["drivers", "Corredores / Pilotos", "firstName, lastName, displayName, number, category, teamId, vehicleId, nationality, license, bloodType, stats"],
        ["vehicles", "Autos / Garaje", "make, model, number, year, category, teamId, powerHp, torqueNm, weightKg, status, vin, transponder, checklist"],
        ["teams", "Equipos / Escuderías", "name, shortName, color, country, principal"],
        ["circuits", "Circuitos", "name, city, country, lengthM, corners, direction, record, recordHolder"],
        ["championships", "Campeonatos", "name, season, category, rounds, pointsScale"],
        ["events", "Grandes Premios / Rondas", "name, round, date, circuitId, championshipId, sessions (prácticas, qualy, carrera)"],
        ["standings", "Clasificación Pilotos", "driverId, teamId, position, points, diff, wins, podiums, trend"],
        ["team_standings", "Clasificación Equipos", "teamId, position, points, diff, wins, podiums"]
    ]
    create_styled_table(doc, db_headers, db_data, [1.5, 2.0, 3.0])

    add_callout(
        doc,
        "Ventajas de la Arquitectura Documental en MySQL",
        "Al almacenar el documento JSON validado en el motor relacional, se evitan costosos JOINs multirregistro para subdocumentos como "
        "el historial de resultados por ronda, documentos federativos o listas de chequeo técnico, manteniendo a la vez la "
        "integridad transaccional ACID de MySQL.",
        "info"
    )

    # -------------------------------------------------------------
    # SECCIÓN 5: GUÍA OPERATIVA CRUD: CORREDORES (PILOTOS)
    # -------------------------------------------------------------
    h1 = doc.add_heading("5. Módulo de Corredores (Pilotos): Operaciones CRUD", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El módulo de pilotos permite la administración total de la parrilla deportiva desde la ruta /pilotos y /pilotos/:driverId:"
    )

    doc.add_heading("5.1. Insertar / Inscribir Nuevo Corredor", level=2)
    doc.add_paragraph(
        "Para ingresar un nuevo piloto al campeonato:\n"
        "1. Navega a la sección Pilotos en el menú lateral de la aplicación.\n"
        "2. Haz clic en el botón '+ Inscribir Piloto' ubicado en el encabezado superior.\n"
        "3. Se desplegará el modal interactivo con los siguientes campos estructurados:\n"
        "   • Identificación Deportiva: Nombre, Apellido, Nombre Deportivo (ej. A. Rossi), Dorsal/Número oficial.\n"
        "   • Categoría: Selección entre GT3 Pro, GT3 Pro-Am o GT4.\n"
        "   • Asignación: Escudería perteneciente y Chasis/Vehículo asignado de los existentes en el garaje.\n"
        "   • Biometría y Nacionalidad: País, Código ISO (ej. ESP, MEX, BRA), Bandera emoji, Ciudad natal y Fecha de nacimiento.\n"
        "   • Licencia Federativa: Código oficial FIA, Grado (Platinum, Gold, Silver, Bronze) y fecha de expiración.\n"
        "   • Ficha Médica de Seguridad: Grupo sanguíneo (A+, O+, etc.), vigencia de aptitud médica y contacto de emergencia.\n"
        "   • Perfil Físico y Biografía: Altura en cm, peso en kg, diseño del casco y reseña biográfica deportiva.\n"
        "4. Presiona 'Inscribir Piloto'. El sistema generará su ID único, lo añadirá a la tabla de posiciones con 0 puntos, "
        "actualizará el contador de pilotos y enviará el registro a MySQL."
    )

    doc.add_heading("5.2. Modificar / Actualizar un Corredor", level=2)
    doc.add_paragraph(
        "Para editar los datos de un piloto ya registrado:\n"
        "• Opción A (Desde el listado): En la tarjeta del piloto, haz clic en el icono del lápiz (Editar) situado en la esquina superior derecha.\n"
        "• Opción B (Desde el detalle): Ingresa a la ficha individual del corredor y haz clic en el botón 'Editar piloto' del encabezado.\n"
        "El modal se abrirá precargado con la información actual. Realiza las modificaciones deseadas y presiona 'Actualizar Piloto'. "
        "Las vistas se actualizarán instantáneamente y se disparará una notificación de confirmación."
    )

    doc.add_heading("5.3. Eliminar / Retirar un Corredor", level=2)
    doc.add_paragraph(
        "Para dar de baja a un corredor de la competición:\n"
        "1. Haz clic en el icono de papelera (Eliminar) en la tarjeta del piloto, o en el botón rojo 'Eliminar' de su página de detalle.\n"
        "2. Se abrirá la ventana modal de confirmación ConfirmDeleteModal advirtiendo que la acción es irreversible.\n"
        "3. Confirma haciendo clic en 'Eliminar definitivamente'.\n"
        "4. El sistema removerá al piloto del listado, suprimirá su fila en la tabla de clasificaciones, recalculará el ranking y "
        "emitirá una petición DELETE /api/drivers/:id y DELETE /api/standings/:id hacia MySQL."
    )

    # -------------------------------------------------------------
    # SECCIÓN 6: GUÍA OPERATIVA CRUD: AUTOS (VEHÍCULOS / GARAJE)
    # -------------------------------------------------------------
    h1 = doc.add_heading("6. Módulo de Garaje (Autos): Operaciones CRUD", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El módulo de parque cerrado y vehículos gestiona la flota técnica en las rutas /garaje y /garaje/:vehicleId:"
    )

    doc.add_heading("6.1. Registrar Nuevo Vehículo de Competición", level=2)
    doc.add_paragraph(
        "Para dar de alta un nuevo auto en el parque cerrado:\n"
        "1. Ingresa a la sección Garaje del menú de navegación.\n"
        "2. Presiona el botón '+ Registrar Vehículo' en la cabecera.\n"
        "3. Completa los datos requeridos en el modal técnico:\n"
        "   • Chasis: Marca (ej. Ferrari, Porsche, BMW), Modelo (ej. 296 GT3, 911 GT3 R), Número dorsal y Año.\n"
        "   • Equipo y Categoría: Escudería propietaria y categoría técnica.\n"
        "   • Estado Scrutineering: Aprobado (listo para pista), Pendiente, Observado o Rechazado.\n"
        "   • Motor y Rendimiento: Tipo de motor (ej. 3.0L V6 Twin-Turbo), Cilindrada, Potencia en HP, Par en Nm y Color distintivo.\n"
        "   • Chasis y Peso: Peso real en kg y peso mínimo homologado por Balance of Performance (BoP).\n"
        "   • Transmisión y Neumáticos: Tipo de caja secuencial, tracción (RWD/AWD), capacidad de tanque y especificación de gomas.\n"
        "   • Fiscalización: Número de bastidor (VIN), código de homologación FIA, código de transponder oficial y notas del comisario.\n"
        "4. Haz clic en 'Registrar Vehículo'. El chasis quedará disponible de inmediato para ser asignado a cualquier corredor."
    )

    doc.add_heading("6.2. Modificar Datos Técnicos del Vehículo", level=2)
    doc.add_paragraph(
        "Puedes actualizar las especificaciones mecánicas o el estado de verificación técnica:\n"
        "• Haz clic en el botón 'Editar' en la tarjeta superior del auto o en el botón 'Editar vehículo' de la página de detalle.\n"
        "• Ajusta los parámetros (por ejemplo, cambio de lastre BoP, potencia modificada o actualización de estado de Pendiente a Aprobado).\n"
        "• Guarda los cambios para que se reflejen en los resúmenes de flota y en la base de datos."
    )

    doc.add_heading("6.3. Eliminar / Retirar Vehículo de Flota", level=2)
    doc.add_paragraph(
        "Si un chasis es retirado del campeonato o sufre daños irreparables:\n"
        "• Presiona el botón 'Eliminar' en la tarjeta o ficha técnica.\n"
        "• El diálogo de seguridad solicitará la confirmación del retiro.\n"
        "• Tras confirmar, el vehículo es eliminado del estado local y de MySQL mediante DELETE /api/vehicles/:id."
    )

    # -------------------------------------------------------------
    # SECCIÓN 7: REFERENCIA DE ENDPOINTS DE LA API REST
    # -------------------------------------------------------------
    h1 = doc.add_heading("7. Especificación de la API RESTful (Backend Express)", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "El servidor backend Express expone un conjunto de endpoints REST estándar sobre el puerto 3001, "
        "con proxy transparente configurado en Vite en la ruta /api:"
    )

    api_headers = ["Método", "Ruta del Endpoint", "Parámetros / Body", "Descripción y Respuesta"]
    api_data = [
        ["GET", "/api/health", "Ninguno", "Comprueba conectividad con MySQL. Retorna {\"status\":\"ok\"} (200) o error 500."],
        ["GET", "/api/counts", "Ninguno", "Retorna el conteo total de registros en cada una de las 8 tablas de la base de datos."],
        ["GET", "/api/:resource", "resource (drivers, vehicles, teams...)", "Retorna un arreglo JSON con todos los registros del catálogo ordenados por created_at."],
        ["GET", "/api/:resource/:id", "resource, id (en URL)", "Obtiene un registro específico por su clave primaria. Retorna 404 si no existe."],
        ["POST", "/api/:resource", "JSON con campos requeridos", "Inserta un nuevo registro. Valida tipos obligatorios. Retorna 201 Created."],
        ["PUT", "/api/:resource/:id", "JSON con modificaciones", "Actualiza un registro existente por id. Retorna 200 con el objeto modificado."],
        ["DELETE", "/api/:resource/:id", "resource, id (en URL)", "Elimina permanentemente el registro con ese ID de la tabla correspondiente en MySQL."],
        ["POST", "/api/bootstrap", "JSON con catálogos iniciales", "Carga datos de demostración masivos si la base de datos está completamente vacía."]
    ]
    create_styled_table(doc, api_headers, api_data, [0.9, 1.8, 1.8, 2.0])

    doc.add_heading("7.1. Ejemplos de Cargas Útiles (Payloads JSON)", level=2)
    doc.add_paragraph(
        "Ejemplo de Payload para creación de Piloto (POST /api/drivers):\n"
        "{\n"
        "  \"firstName\": \"Carlos\",\n"
        "  \"lastName\": \"Sainz\",\n"
        "  \"displayName\": \"C. Sainz\",\n"
        "  \"number\": 55,\n"
        "  \"category\": \"GT3 Pro\",\n"
        "  \"teamId\": \"t-aurora\",\n"
        "  \"vehicleId\": \"v-aurora-1\",\n"
        "  \"nationality\": \"España\",\n"
        "  \"countryCode\": \"ESP\",\n"
        "  \"flag\": \"🇪🇸\"\n"
        "}"
    )

    doc.add_paragraph(
        "Ejemplo de Payload para creación de Vehículo (POST /api/vehicles):\n"
        "{\n"
        "  \"make\": \"Aston Martin\",\n"
        "  \"model\": \"Vantage AMR GT3\",\n"
        "  \"number\": 77,\n"
        "  \"year\": 2026,\n"
        "  \"teamId\": \"t-aurora\",\n"
        "  \"category\": \"GT3 Pro\",\n"
        "  \"powerHp\": 545,\n"
        "  \"weightKg\": 1260,\n"
        "  \"status\": \"approved\"\n"
        "}"
    )

    # -------------------------------------------------------------
    # SECCIÓN 8: MANUAL DE INSTALACIÓN Y PUESTA EN MARCHA
    # -------------------------------------------------------------
    h1 = doc.add_heading("8. Manual de Puesta en Marcha y Configuración", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Para ejecutar la solución en un entorno local de desarrollo, sigue los pasos indicados:"
    )

    doc.add_heading("Paso 1: Configurar la Base de Datos MySQL (Opcional pero recomendado)", level=2)
    doc.add_paragraph(
        "1. Asegúrate de tener una instancia de MySQL en ejecución (local o en contenedor).\n"
        "2. Ejecuta el script database/schema.sql en tu cliente MySQL para crear la base de datos apex_race_control y sus tablas.\n"
        "3. Crea un archivo .env en la raíz del proyecto copiando .env.example y ajusta las credenciales:\n"
        "   DB_HOST=127.0.0.1\n"
        "   DB_PORT=40628 (o 3306 según tu configuración)\n"
        "   DB_USER=root\n"
        "   DB_PASSWORD=tu_contraseña\n"
        "   DB_DATABASE=apex_race_control\n"
        "   API_PORT=3001"
    )

    doc.add_heading("Paso 2: Iniciar el Servidor de la API", level=2)
    doc.add_paragraph(
        "En una terminal de comandos, ejecuta:\n"
        "npm run api\n"
        "El servidor confirmará en consola: 'API lista en http://127.0.0.1:3001'."
    )

    doc.add_heading("Paso 3: Iniciar la Aplicación Frontend Vite", level=2)
    doc.add_paragraph(
        "En una segunda terminal, ejecuta:\n"
        "npm run dev\n"
        "Abre en tu navegador la dirección indicada, típicamente http://localhost:5173.\n"
        "Inicia sesión con el rol de Organizador o Dirección de Carrera para disponer de todos los permisos operativos."
    )

    add_callout(
        doc,
        "Modo Independiente (Offline First)",
        "Si no dispones de un servidor MySQL activo en tu máquina, la aplicación funcionará al 100% en modo LocalStorage. "
        "Podrás crear, modificar y eliminar pilotos y autos sin ningún tipo de error ni bloqueo en pantalla.",
        "success"
    )

    # -------------------------------------------------------------
    # SECCIÓN 9: PREGUNTAS FRECUENTES Y SOLUCIÓN DE PROBLEMAS
    # -------------------------------------------------------------
    h1 = doc.add_heading("9. Preguntas Frecuentes y Resolución de Problemas", level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)

    faq_data = [
        ("¿Qué sucede si elimino a un corredor que tenía puntos en el campeonato?",
         "La fila de clasificación correspondiente a ese piloto se elimina automáticamente de la tabla de standings y el sistema reordena las posiciones de los pilotos restantes para mantener la coherencia numérica del ranking."),
        ("¿Puedo volver a restaurar los datos iniciales de fábrica?",
         "Sí. El store useCompetition cuenta con la función resetToDefaults(), y en la sección Administración → Base de Datos se incluye el botón 'Cargar datos de ejemplo' para poblar de nuevo MySQL."),
        ("¿Cómo se refleja un nuevo piloto en el Live Timing?",
         "La capa de datos src/lib/data.ts consulta dinámicamente el store reactivo useCompetition. Por ello, cualquier piloto recién creado o renombrado es reconocido inmediatamente por el motor de cronometraje y simulación."),
        ("¿Qué validaciones se aplican al insertar un auto?",
         "El sistema verifica que la marca, el modelo y el dorsal sean válidos, que la potencia y el peso sean valores numéricos positivos, y asigna automáticamente valores por defecto a los documentos federativos y listas de chequeo técnico.")
    ]

    for q, a in faq_data:
        p_q = doc.add_paragraph()
        r_q = p_q.add_run(f"• {q}\n")
        r_q.bold = True
        r_q.font.color.rgb = RGBColor(15, 23, 42)
        r_a = p_q.add_run(a)
        r_a.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(output_path)
    print(f"Documento generado exitosamente en: {output_path}")

if __name__ == '__main__':
    target = os.path.abspath("Documentacion_Apex_Race_Control.docx")
    build_word_documentation(target)
